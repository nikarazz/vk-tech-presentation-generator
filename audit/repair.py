"""Исправление найденных проблем."""
from __future__ import annotations

from shared.schemas.audit import AuditReport, IssueType


def repair(presentation, audit_report: AuditReport, design_system):
    """Применить исправления по отчёту аудита."""
    slide_w = design_system.slide_width
    slide_h = design_system.slide_height

    for issue in audit_report.issues:
        if issue.slide >= len(presentation.slides):
            continue
        slide = presentation.slides[issue.slide]
        if issue.type == IssueType.BOUNDS:
            _fix_bounds(slide, issue, slide_w, slide_h)
        elif issue.type == IssueType.OVERLAP:
            _fix_overlap(slide, issue)
        elif issue.type == IssueType.TYPOGRAPHY:
            _fix_typography(slide, issue)
        elif issue.type == IssueType.COLORS:
            _fix_colors(slide, issue)

    return presentation


def _find_shape(slide, shape_id):
    for shape in slide.shapes:
        if getattr(shape, "shape_id", None) == shape_id:
            return shape
    return None


def _fix_bounds(slide, issue, slide_w, slide_h):
    shape = _find_shape(slide, issue.shape_id) if issue.shape_id else None
    if shape is None:
        return
    if shape.left is None or shape.left < 0:
        shape.left = 0
    if shape.top is None or shape.top < 0:
        shape.top = 0
    if shape.left + shape.width > slide_w:
        shape.left = max(0, slide_w - shape.width)
    if shape.top + shape.height > slide_h:
        shape.top = max(0, slide_h - shape.height)


def _fix_overlap(slide, issue):
    shape = _find_shape(slide, issue.shape_id) if issue.shape_id else None
    if shape is None:
        return
    shape.top = (shape.top or 0) + 20


def _fix_typography(slide, issue):
    from pptx.util import Pt
    shape = _find_shape(slide, issue.shape_id) if issue.shape_id else None
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None and run.font.size.pt < 12:
                run.font.size = Pt(12)


def _fix_colors(slide, issue):
    from pptx.dml.color import RGBColor
    shape = _find_shape(slide, issue.shape_id) if issue.shape_id else None
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            try:
                run.font.color.rgb = RGBColor(0, 0, 0)
            except Exception:
                pass
