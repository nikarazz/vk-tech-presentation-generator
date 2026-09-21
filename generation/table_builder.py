from pptx.util import Emu, Pt


def build_table(slide, spec: dict, ds: dict, bbox: dict):
    """
    spec = {
        "headers": ["Метрика", "Q1", "Q2", "Q3", "Q4"],
        "rows": [
            ["DAU", "10M", "12M", "15M", "18M"],
            ["Retention", "40%", "42%", "45%", "48%"],
        ],
    }
    """
    headers = spec["headers"]
    rows = spec["rows"]
    
    n_rows = len(rows) + 1  # +1 для заголовка
    n_cols = len(headers)
    
    table_shape = slide.shapes.add_table(
        n_rows, n_cols,
        Emu(bbox["x_emu"]),
        Emu(bbox["y_emu"]),
        Emu(bbox["w_emu"]),
        Emu(bbox["h_emu"]),
    )
    table = table_shape.table
    
    # Заголовки
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = str(header)
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(14)
    
    # Данные
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(12)
    
    return table