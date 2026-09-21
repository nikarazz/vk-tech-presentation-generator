from pptx import Presentation
from pathlib import Path


def _emu_to_px(emu: int) -> int:
    """EMU → пиксели (примерно)."""
    return int(emu / 9525)


def export_html(pptx_path: str, output_path: str) -> None:
    """
    Конвертирует .pptx → HTML.
    Каждый слайд — отдельная секция. Текст, таблицы, координаты сохраняются.
    """
    prs = Presentation(pptx_path)
    
    slide_width_px = _emu_to_px(prs.slide_width)
    slide_height_px = _emu_to_px(prs.slide_height)
    
    html_parts = [
        "<!DOCTYPE html>",
        "<html lang='ru'>",
        "<head>",
        "<meta charset='UTF-8'>",
        "<title>Presentation</title>",
        "<style>",
        "body { font-family: Arial, sans-serif; background: #f0f0f0; margin: 0; padding: 20px; }",
        f".slide {{ position: relative; width: {slide_width_px}px; height: {slide_height_px}px; background: #fff; margin: 20px auto; box-shadow: 0 2px 8px rgba(0,0,0,0.1); overflow: hidden; }}",
        ".el { position: absolute; }",
        ".el-text { white-space: pre-wrap; }",
        "table { border-collapse: collapse; width: 100%; height: 100%; }",
        "td, th { border: 1px solid #ccc; padding: 6px 10px; font-size: 14px; }",
        "th { background: #4a7ebb; color: #fff; font-weight: bold; }",
        "</style>",
        "</head>",
        "<body>",
    ]
    
    for i, slide in enumerate(prs.slides, start=1):
        html_parts.append(f"<div class='slide' data-slide='{i}'>")
        
        for shape in slide.shapes:
            # Таблицы
            if shape.has_table:
                html_parts.append(
                    f"<div class='el' style='left:{_emu_to_px(shape.left)}px; "
                    f"top:{_emu_to_px(shape.top)}px; "
                    f"width:{_emu_to_px(shape.width)}px; "
                    f"height:{_emu_to_px(shape.height)}px;'>"
                )
                html_parts.append("<table>")
                for row in shape.table.rows:
                    html_parts.append("<tr>")
                    for cell in row.cells:
                        html_parts.append(f"<td>{cell.text}</td>")
                    html_parts.append("</tr>")
                html_parts.append("</table>")
                html_parts.append("</div>")
                continue
            
            # Текст
            if not shape.has_text_frame:
                continue
            
            text = shape.text_frame.text
            if not text.strip():
                continue
            
            font_size = 16
            bold = False
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size:
                        font_size = int(run.font.size.pt)
                    if run.font.bold:
                        bold = True
                    break
                break
            
            style = (
                f"left:{_emu_to_px(shape.left)}px; "
                f"top:{_emu_to_px(shape.top)}px; "
                f"width:{_emu_to_px(shape.width)}px; "
                f"height:{_emu_to_px(shape.height)}px; "
                f"font-size:{font_size}px; "
                f"{'font-weight:bold;' if bold else ''}"
            )
            
            text_html = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            html_parts.append(f"<div class='el el-text' style='{style}'>{text_html}</div>")
        
        html_parts.append("</div>")
    
    html_parts.extend(["</body>", "</html>"])
    
    Path(output_path).write_text("\n".join(html_parts), encoding="utf-8")
    print(f"[B] Exported HTML: {output_path}")