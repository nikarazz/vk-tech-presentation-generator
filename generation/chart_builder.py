from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu


CHART_TYPE_MAP = {
    "bar":    XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line":   XL_CHART_TYPE.LINE_MARKERS,
    "pie":    XL_CHART_TYPE.PIE,
    "donut":  XL_CHART_TYPE.DOUGHNUT,
    "area":   XL_CHART_TYPE.AREA,
    "stacked_bar": XL_CHART_TYPE.COLUMN_STACKED,
}


def build_chart(slide, spec: dict, ds: dict, bbox: dict):
    """Создаёт график на слайде."""
    chart_type = CHART_TYPE_MAP.get(
        spec.get("type", "bar"),
        XL_CHART_TYPE.COLUMN_CLUSTERED,
    )

    chart_data = CategoryChartData()
    chart_data.categories = spec["categories"]

    for series in spec["series"]:
        chart_data.add_series(series["name"], series["values"])

    chart_shape = slide.shapes.add_chart(
        chart_type,
        Emu(bbox["x_emu"]),
        Emu(bbox["y_emu"]),
        Emu(bbox["w_emu"]),
        Emu(bbox["h_emu"]),
        chart_data,
    )

    chart = chart_shape.chart

    # Заголовок
    if spec.get("title"):
        chart.has_title = False
        chart.chart_title.text_frame.text = spec["title"]
    else:
        chart.has_title = False

    # Легенда
    is_pie = chart_type in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
    if is_pie or len(spec["series"]) > 1:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    else:
        chart.has_legend = False

    # Подписи осей
    if not is_pie:
        try:
            chart.category_axis.has_title = True
            chart.category_axis.axis_title.text_frame.text = spec.get("x_label", "")
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = spec.get("y_label", "")
        except Exception:
            pass

    return chart
